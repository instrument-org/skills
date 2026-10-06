#!/usr/bin/env python3
"""Write an HTML page that draws every slide of a deck, for the agent's browser to open, print and look at.

The page draws slides with the same renderer Instrument uses to show a deck to the person, loaded from a CDN, so it needs a network connection. Each slide prints on its own page at the deck's own size. When every slide is drawn, the page sets `document.body.dataset.ready` to the slide count; on failure it sets `document.body.dataset.error`.
"""

import argparse
import base64
import html
import json
import sys
from pathlib import Path

# Printed in place of the slides when the page is printed before they are
# drawn, and refused by thumbnail.py, so an early print cannot pass as a deck.
EARLY_PRINT = "This preview was printed before its slides were drawn. Wait for document.body.dataset.ready, then print again."

# The renderer and the version Instrument itself ships, so the preview draws
# what the person will see when they open the deck there.
RENDERER = "@extend-ai/react-pptx@0.1.2"
REACT = "19.2.0"
EMU_PER_PIXEL = 9525

PAGE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>__TITLE__</title>
<link rel="stylesheet" href="https://cdn.jsdelivr.net/npm/__RENDERER__/dist/index.css">
<style>
  @page { size: __W__px __H__px; margin: 0; }
  html, body { margin: 0; background: #fff; }
  #pages > div { width: __W__px; height: __H__px; overflow: hidden; break-after: page; }
  #pages > div:last-child { break-after: auto; }
  body[data-ready] #viewer { display: none; }
  @media print {
    body:not([data-ready]) > * { display: none !important; }
    body:not([data-ready])::before { content: "__EARLY__"; display: block; padding: 48px; font: 28px sans-serif; }
  }
</style>
</head>
<body>
<div id="viewer"></div>
<div id="pages"></div>
<script type="application/octet-stream" id="deck">__DECK__</script>
<script type="module">
import React from "https://esm.sh/react@__REACT__";
import { createRoot } from "https://esm.sh/react-dom@__REACT__/client";
import { ReactPptxViewer, setWasmSource } from "https://esm.sh/__RENDERER__?deps=react@__REACT__,react-dom@__REACT__";

const fail = (error) => { document.body.dataset.error = String(error?.message ?? error); };
try {
  // Handing the renderer a Response parses on this page's own thread; its
  // worker would be a cross-origin script, which a file:// page cannot start.
  setWasmSource(await fetch("https://cdn.jsdelivr.net/npm/__RENDERER__/dist/pptx_wasm_bg.wasm"));
  const bytes = Uint8Array.from(atob(document.getElementById("deck").textContent.trim()), (c) => c.charCodeAt(0));
  const drawn = new Map();
  let total = 0;
  const finish = async () => {
    await document.fonts.ready;
    const pages = document.getElementById("pages");
    for (let index = 0; index < total; index++) {
      const page = document.createElement("div");
      // Moved, not copied: a copied canvas (a chart) loses its pixels.
      page.append(drawn.get(index));
      pages.append(page);
    }
    document.body.dataset.ready = String(total);
  };
  createRoot(document.getElementById("viewer")).render(
    React.createElement(ReactPptxViewer, {
      source: bytes,
      mode: "continuous",
      virtualization: false,
      zoom: 100,
      height: "auto",
      showToolbar: false,
      showThumbnails: false,
      viewportStyle: { height: "auto", overflow: "visible", padding: 0 },
      onLoad: (presentation) => { total = presentation.document.slides.length; },
      onSlideRendered: (index, element) => {
        drawn.set(index, element);
        if (total > 0 && drawn.size === total) requestAnimationFrame(() => finish().catch(fail));
      },
      onError: fail,
    }),
  );
} catch (error) {
  fail(error);
}
</script>
</body>
</html>
"""


def main():
    parser = argparse.ArgumentParser(
        description="Write an HTML page that draws every slide of a .pptx for the browser to print and look at"
    )
    parser.add_argument("input", help="Input .pptx file")
    parser.add_argument("--output", help="HTML file to write (default: <input>.preview.html beside the deck)")
    args = parser.parse_args()

    try:
        from pptx import Presentation
    except ImportError:
        sys.exit("python-pptx is unavailable. Reload this skill to retry dependency setup.")

    deck = Path(args.input)
    if not deck.is_file():
        sys.exit(f"No such file: {deck}")
    prs = Presentation(str(deck))
    width = round(prs.slide_width / EMU_PER_PIXEL)
    height = round(prs.slide_height / EMU_PER_PIXEL)
    output = Path(args.output) if args.output else deck.with_name(f"{deck.stem}.preview.html")

    page = (
        PAGE.replace("__TITLE__", html.escape(f"{deck.name} preview"))
        .replace("__EARLY__", EARLY_PRINT)
        .replace("__RENDERER__", RENDERER)
        .replace("__REACT__", REACT)
        .replace("__W__", str(width))
        .replace("__H__", str(height))
        .replace("__DECK__", base64.b64encode(deck.read_bytes()).decode("ascii"))
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(page, encoding="utf-8")
    print(json.dumps({"preview": str(output), "slides": len(prs.slides), "slideSize": [width, height]}))


if __name__ == "__main__":
    main()
